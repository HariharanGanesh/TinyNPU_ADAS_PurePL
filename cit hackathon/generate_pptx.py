from pptx import Presentation
from pptx.util import Inches, Pt
import os

def add_slide(prs, title, content):
    slide_layout = prs.slide_layouts[1] # Title and Content layout
    slide = prs.slides.add_slide(slide_layout)
    
    title_shape = slide.shapes.title
    title_shape.text = title
    
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    
    # Enable word wrap
    body_shape.text_frame.word_wrap = True
    
    # Split content by newlines
    lines = content.split('\n')
    for i, line in enumerate(lines):
        line = line.strip()
        if not line:
            continue
            
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
            
        # Basic bullet handling
        if line.startswith('* '):
            p.text = line[2:]
            p.level = 0
        elif line.startswith('    * '):
            p.text = line[6:]
            p.level = 1
        else:
            p.text = line
            p.level = 0

def main():
    prs = Presentation()
    
    # Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "TinyNPU - Board-Agnostic Edge AI Accelerator"
    subtitle.text = "DEEPSPRINT Hackathon\nDomain: VLSI & Semiconductor Tech"
    
    # Slide 1: Team Details
    add_slide(prs, "Team Details",
              "* a. Team name: [Your Team Name]\n"
              "* b. Team leader name: [Your Name]\n"
              "* c. Team size: [E.g., 1 / 3 / 4]")
              
    # Slide 2: Problem
    add_slide(prs, "Problem",
              "* 1. Problem Statement: Deploying AI at the edge faces a critical bottleneck: hardware is either too power-hungry or rigidly tied to specific vendor boards. Current tech suffers from high latency jitter, massive power draw, and complex software-stack dependencies (OS overhead), all of which our 100% PL-based TinyNPU eliminates.\n"
              "* 2. Who has the problem?: Defense contractors, drone manufacturers, and aerospace engineers who require deterministic, real-time, and low-power AI inference (like target detection) in extreme, resource-constrained environments.\n"
              "* 3. Current solution?:\n"
              "    * CPUs/GPUs: Consume too much power, are physically bulky, and suffer from latency jitter due to Operating System overhead.\n"
              "    * Vendor-Locked NPUs: Rigidly tied to specific SoC architectures (like the Zynq Processing System). They require complex ARM software stacks to boot, making them impossible to easily synthesize as standalone IP on alternative FPGA fabrics.\n"
              "* 4. Why important?: In defense and aerospace, deterministic execution (zero-jitter) and ultra-low power are non-negotiable. A board-agnostic, zero-OS overhead accelerator drastically reduces time-to-market and brings AI to edge environments where traditional hardware fails.")

    # Slide 3: Solution - Describe your Innovation
    add_slide(prs, "Solution - Describe your Innovation",
              "TinyNPU: A 100% PL-Based, Board-Agnostic Neural Processing Unit\n"
              "* Fully Synthesizable IP: Unlike traditional NPUs that rely on specific Processing Systems (like Zynq PS), our TinyNPU is designed 100% in Programmable Logic (PL). It can be ported to any FPGA fabric seamlessly.\n"
              "* High-Efficiency Architecture: Custom-designed silicon architecture optimized specifically for low-latency, low-power inference at the edge.\n"
              "* Real-World Applicability: Capable of running quantized neural networks (e.g., MobileNet/TinyYOLO) for real-time target detection directly on the hardware fabric.")

    # Slide 4: Customer
    add_slide(prs, "Customer",
              "* 1. Target Customer: B2B Defense technology firms, Aerospace agencies (ISRO/DRDO vendors), Autonomous robotics startups, and semiconductor IP integrators.\n"
              "* 2. Market Size: The Edge AI Hardware market is projected to reach $38.8 Billion by 2030 (CAGR of 18.8%). Our initial focus is the rapidly growing edge defense and robotics niche, valued at over $5 Billion.")

    # Slide 5: Business
    add_slide(prs, "Business",
              "* 1. Revenue model:\n"
              "    * IP Licensing: Licensing the TinyNPU RTL/bitstream to semiconductor firms and defense contractors for integration into their custom SoCs.\n"
              "    * NRE (Non-Recurring Engineering): Customizing the NPU architecture for client-specific neural network workloads.\n"
              "* 2. Selling price:\n"
              "    * Base IP License: $50,000 - $100,000 + royalties per chip.\n"
              "    * Note: This is a highly competitive B2B semiconductor IP pricing model, significantly cheaper than developing custom silicon in-house.")

    # Slide 6: Validation
    add_slide(prs, "Validation",
              "* 1. Prototype? (Yes or No): Yes. (We have the RTL/bitstream ready and validated on FPGA fabric, benchmarking power and latency).\n"
              "* 2. TRL (Technology Readiness Level): TRL 4 (Component and/or breadboard validation in laboratory environment). We are actively moving towards TRL 5 (System validation).")

    # Slide 7: IP (Intellectual Property)
    add_slide(prs, "IP (Intellectual Property)",
              "* 1. Novelty: The hardware-software interface is entirely decoupled from hard processors. The custom routing and scheduling within the 100% PL-based architecture achieves a unique balance of ultra-low latency and minimal area footprint not found in off-the-shelf IPs.\n"
              "* 2. Patent? (Yes or No): No (Currently maintaining as trade secret / planning to file provisional patent before commercial licensing).")

    prs.save('DEEPSPRINT_Pitch.pptx')
    print("Successfully generated DEEPSPRINT_Pitch.pptx")

if __name__ == '__main__':
    main()
